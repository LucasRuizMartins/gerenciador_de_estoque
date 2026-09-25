"""Formatação visual padronizada para planilhas Excel (openpyxl).

Exporta uma única função pública:
    formatar_aba(ws, df)

Aplica:
  - Cabeçalho: fundo azul escuro (#1F4E79), texto branco, negrito, centralizado.
  - Linhas de dados: alternância azul claro / branco (efeito zebra).
  - Bordas finas em azul médio em todas as células.
  - Largura de coluna ajustada ao conteúdo (mín 10, máx 45 caracteres).
  - Cabeçalho congelado na linha 2 (freeze_panes = "A2").
  - Formato numérico automático detectado pelo nome da coluna:
      * palavras-chave de valor/moeda → #,##0.00
      * palavras-chave de percentual   → 0.00%
      * palavras-chave de quantidade   → #,##0
      * demais colunas                 → texto (alinhado à esquerda)
"""

from __future__ import annotations

import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Paleta de cores ────────────────────────────────────────────────────────────
_COR_HEADER_BG = "1F4E79"   # azul escuro
_COR_HEADER_FT = "FFFFFF"   # branco
_COR_LINHA_PAR = "DCE6F1"   # azul claro (linhas pares)
_COR_LINHA_IMP = "FFFFFF"   # branco (linhas ímpares)
_COR_BORDA     = "B8CCE4"   # azul médio

# ── Formatos numéricos ─────────────────────────────────────────────────────────
_FMT_MOEDA = "#,##0.00"
_FMT_PCT   = "0.00%"
_FMT_INT   = "#,##0"

# Palavras-chave para detecção automática de formato por nome de coluna
_KEYWORDS_MOEDA = ("vl.", "valor", "compra", "vencimento", "aquisição", "aquisicao",
                   "pago", "nominal", "presente", "ticket", "pdd")
_KEYWORDS_PCT   = ("deságio", "desagio", "retorno", "taxa", "%", "pct", "mediana")
_KEYWORDS_INT   = ("qtd", "qtde", "quantidade", "titulos", "títulos",
                   "sacados", "cedentes", "dias", "prazo")


def _detectar_formato(nome_coluna: str) -> str | None:
    """Retorna o formato numérico Excel adequado para a coluna, ou None para texto."""
    n = nome_coluna.lower()
    if any(k in n for k in _KEYWORDS_MOEDA):
        return _FMT_MOEDA
    if any(k in n for k in _KEYWORDS_PCT):
        return _FMT_PCT
    if any(k in n for k in _KEYWORDS_INT):
        return _FMT_INT
    return None


def formatar_aba(ws, df: pd.DataFrame) -> None:
    """Aplica formatação visual e de tipo a uma aba do openpyxl.

    Args:
        ws:  Worksheet do openpyxl (writer.sheets["Nome da Aba"]).
        df:  DataFrame correspondente (usado para leitura dos nomes de colunas).
    """
    thin_side   = Side(style="thin", color=_COR_BORDA)
    borda_fina  = Border(left=thin_side, right=thin_side,
                         top=thin_side, bottom=thin_side)

    fill_header = PatternFill("solid", fgColor=_COR_HEADER_BG)
    fill_par    = PatternFill("solid", fgColor=_COR_LINHA_PAR)
    fill_imp    = PatternFill("solid", fgColor=_COR_LINHA_IMP)

    font_header  = Font(bold=True, color=_COR_HEADER_FT, name="Calibri", size=11)
    font_data    = Font(name="Calibri", size=10)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left   = Alignment(horizontal="left",   vertical="center")
    align_right  = Alignment(horizontal="right",  vertical="center")

    # Mapa coluna_idx (1-based) → formato numérico
    col_fmts = {i + 1: _detectar_formato(c) for i, c in enumerate(df.columns)}

    # ── Cabeçalho (linha 1) ────────────────────────────────────────────
    for cell in ws[1]:
        cell.fill      = fill_header
        cell.font      = font_header
        cell.alignment = align_center
        cell.border    = borda_fina
    ws.row_dimensions[1].height = 20

    # ── Linhas de dados ────────────────────────────────────────────────
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill_row = fill_par if row_idx % 2 == 0 else fill_imp
        for cell in row:
            cell.fill   = fill_row
            cell.font   = font_data
            cell.border = borda_fina
            fmt = col_fmts.get(cell.column)
            if fmt:
                cell.number_format = fmt
                cell.alignment     = align_right
            else:
                cell.alignment     = align_left
        ws.row_dimensions[row_idx].height = 16

    # ── Largura automática de colunas ──────────────────────────────────
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = max(
            (len(str(cell.value)) if cell.value is not None else 0)
            for cell in col_cells
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 45)

    # ── Congela cabeçalho ──────────────────────────────────────────────
    ws.freeze_panes = "A2"
